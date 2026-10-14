"""CLI 端到端测试。

用 Typer 的 CliRunner 跑真实命令 —— 参数解析、插件发现、
用例编排、报告导出全链路。打出来的报告再用 openpyxl
读回来验证，而不是只看退出码。

**退出码也是契约：** 0=成功；1=有插件不可用；2=输入文件不存在。
脚本化调用（CI、同事的批处理）靠它判断成败。
"""

from __future__ import annotations

import pytest

from conftest import SAMPLE_DATA
from automotive_dev_support.apps.cli.main import app

pytest.importorskip("typer")
from typer.testing import CliRunner  # noqa: E402

pytest.importorskip("openpyxl")
from openpyxl import load_workbook  # noqa: E402

runner = CliRunner()


def _compare_args(**overrides) -> list[str]:
    args = [
        "compare",
        "--baseline-a2l", str(SAMPLE_DATA / "sample_v1.a2l"),
        "--candidate-a2l", str(SAMPLE_DATA / "sample_v2.a2l"),
    ]
    for key, value in overrides.items():
        flag = "--" + key.replace("_", "-")
        if value is None:
            continue
        args.extend([flag, str(value)])
    return args


class TestCompareCommand:
    def test_cdfx_backed_compare(self):
        result = runner.invoke(
            app,
            _compare_args(
                baseline_values=SAMPLE_DATA / "sample_v1.cdfx",
                candidate_values=SAMPLE_DATA / "sample_v2.cdfx",
            ),
        )
        assert result.exit_code == 0, result.output
        assert "新增 1 / 删除 1 / 改动 3" in result.output
        assert "Fuel_Cut_Threshold" in result.output
        # 未挂值的量会告警，但不中断
        assert "⚠" in result.output

    def test_hex_backed_compare(self):
        result = runner.invoke(
            app,
            _compare_args(
                baseline_image=SAMPLE_DATA / "sample_v1.hex",
                candidate_image=SAMPLE_DATA / "sample_v2.hex",
            ),
        )
        assert result.exit_code == 0, result.output
        assert "新增 1 / 删除 1 / 改动 3" in result.output
        # 镜像提供了 Torque_Map 的值，所以它不该出现在缺失告警里
        assert "Torque_Map" not in result.output.split("需要关注")[0]

    def test_s19_image_is_accepted(self):
        """.s19 与 .hex 描述同一份内存，比对结果必须一致。"""
        hex_result = runner.invoke(
            app,
            _compare_args(
                baseline_image=SAMPLE_DATA / "sample_v1.hex",
                candidate_image=SAMPLE_DATA / "sample_v2.hex",
            ),
        )
        assert hex_result.exit_code == 0, hex_result.output
        # v2 只有 hex 版，用 v1 的 s19 与 v1 的 hex 对比应报告「无变化」
        same = runner.invoke(
            app,
            [
                "compare",
                "--baseline-a2l", str(SAMPLE_DATA / "sample_v1.a2l"),
                "--candidate-a2l", str(SAMPLE_DATA / "sample_v1.a2l"),
                "--baseline-image", str(SAMPLE_DATA / "sample_v1.hex"),
                "--candidate-image", str(SAMPLE_DATA / "sample_v1.s19"),
            ],
        )
        assert same.exit_code == 0, same.output
        assert "需要关注 0 项" in same.output

    def test_report_is_written_and_valid(self, tmp_path):
        out = tmp_path / "report.xlsx"
        result = runner.invoke(
            app,
            _compare_args(
                baseline_values=SAMPLE_DATA / "sample_v1.cdfx",
                candidate_values=SAMPLE_DATA / "sample_v2.cdfx",
                out=out,
            ),
        )
        assert result.exit_code == 0, result.output
        assert out.exists()
        wb = load_workbook(out)
        assert wb.sheetnames == ["汇总", "改动明细", "地址漂移"]
        summary = wb["汇总"]
        texts = [str(c.value) for row in summary.iter_rows() for c in row if c.value]
        assert any("sample_v2.a2l" in t for t in texts)

    def test_missing_file_exits_2(self):
        result = runner.invoke(
            app,
            _compare_args(baseline_values=SAMPLE_DATA / "no_such_file.cdfx"),
        )
        assert result.exit_code == 2
        assert "不存在" in result.output

    def test_no_values_warns_but_succeeds(self):
        """只给 A2L 也能跑 —— 只比描述，不比数值。"""
        result = runner.invoke(app, _compare_args())
        assert result.exit_code == 0, result.output
        assert "未提供值集或内存镜像" in result.output
        # 无数值可比的 4 项：Fuel_Cut 新增 / Temp_Offset 删除 /
        # Idle_Base 与 Ki_Speed 属性变化（Kp_Speed 只剩地址漂移，不算）
        assert "需要关注 4 项" in result.output


class TestPluginsCommand:
    def test_lists_all_plugins(self):
        result = runner.invoke(app, ["plugins"])
        assert result.exit_code == 0, result.output
        for plugin_id in (
            "parser-a2l",
            "parser-cdfx",
            "parser-hex",
            "parser-s19",
            "diff-calibration",
            "export-excel-diff",
        ):
            assert plugin_id in result.output

    def test_help_works(self):
        result = runner.invoke(app, ["--help"])
        assert result.exit_code == 0
        assert "compare" in result.output
