"""Excel 报告导出测试。

导出的东西是要拿出去给人看的，所以这里验证的不只是
「能生成文件」，而是「生成的文件里内容是对的」——
用 openpyxl 读回来逐格断言。生成的 xlsx 打不开、
或者格子里的内容不对，都比没有导出更糟。
"""

from __future__ import annotations

import pytest

from automotive_dev_support.core.models.calibration import (
    ChangeKind,
    ChangeSet,
    CharacteristicChange,
    FieldChange,
    ValueChange,
)
from automotive_dev_support.core.ports import PluginRegistry
from automotive_dev_support.core.skeleton import JobContext
from automotive_dev_support.plugins.export_excel_diff import ExcelDiffExporter

from conftest import PLUGINS_DIR

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import load_workbook  # noqa: E402

CTX = JobContext(job_id="test")


def _changeset() -> ChangeSet:
    return ChangeSet(
        baseline_label="baseline.hex",
        candidate_label="candidate.hex",
        added=[
            CharacteristicChange(path="Fuel_Cut_Threshold", name="Fuel_Cut_Threshold",
                                 kinds=[ChangeKind.ADDED], group="TorqueGroup")
        ],
        removed=[
            CharacteristicChange(path="Temp_Offset", name="Temp_Offset",
                                 kinds=[ChangeKind.REMOVED])
        ],
        changed=[
            CharacteristicChange(
                path="Kp_Speed",
                name="Kp_Speed",
                kinds=[ChangeKind.VALUES_CHANGED],
                group="TorqueGroup",
                address_before=0x80001234,
                address_after=0x80001244,
                value_changes=[ValueChange(index=0, phys_before=850.0, phys_after=900.0)],
            ),
            CharacteristicChange(
                path="Idle_Base",
                name="Idle_Base",
                kinds=[ChangeKind.ATTRS_CHANGED],
                attribute_changes=[
                    FieldChange(field="compu", before="CM_Pct", after="CM_None"),
                    FieldChange(field="unit", before="%", after=""),
                ],
            ),
            CharacteristicChange(
                path="Ki_Speed",
                name="Ki_Speed",
                kinds=[ChangeKind.ADDRESS_MOVED],
                address_before=0x80001236,
                address_after=0x80001246,
            ),
        ],
        warnings=["描述文件里有 1 个量没有值：Torque_Map"],
    )


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    cs = _changeset()
    target = tmp_path_factory.mktemp("report") / "diff.xlsx"
    result = ExcelDiffExporter().export(cs, target, {}, CTX)
    return load_workbook(result)


class TestExcelExport:
    def test_three_sheets(self, report):
        assert report.sheetnames == ["汇总", "改动明细", "地址漂移"]

    def test_summary_counts(self, report):
        ws = report["汇总"]
        cells = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 12)}
        assert cells["基准版本"] == "baseline.hex"
        assert cells["对比版本"] == "candidate.hex"
        assert cells["新增标定量"] == 1
        assert cells["删除标定量"] == 1
        assert cells["改动标定量"] == 3
        assert cells["其中：纯地址漂移"] == 1
        assert cells["需要关注的改动"] == 4  # 新增/删除/数值/属性各一，纯地址漂移不算

    def test_summary_shows_warnings(self, report):
        ws = report["汇总"]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert any("Torque_Map" in t for t in texts)

    def test_details_rows(self, report):
        ws = report["改动明细"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        by_path = {r[0]: r for r in rows}
        assert set(by_path) == {"Fuel_Cut_Threshold", "Temp_Offset", "Kp_Speed", "Idle_Base", "Ki_Speed"}
        # 地址按 16 进制写，且箭头方向是 前 → 后
        assert by_path["Kp_Speed"][5] == "0x80001234 → 0x80001244"
        # 数值变化按物理值展示
        assert "850 → 900" in by_path["Kp_Speed"][7]
        # 属性变化逐条列出
        assert "compu: CM_Pct → CM_None" in by_path["Idle_Base"][6]
        assert "unit: % → " in by_path["Idle_Base"][6]

    def test_pure_address_move_goes_to_its_own_sheet(self, report):
        ws = report["地址漂移"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert rows == [("Ki_Speed", "0x80001236", "0x80001246")]

    def test_meaningful_changes_sort_before_noise(self, report):
        ws = report["改动明细"]
        paths = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
        assert paths[-1] == "Ki_Speed"  # 纯地址漂移排在最后

    def test_map_elements_are_capped(self, tmp_path):
        """16x8 的 MAP 全量列出来没人看 —— 必须截断并注明总数。"""
        cs = ChangeSet(
            changed=[
                CharacteristicChange(
                    path="Torque_Map",
                    kinds=[ChangeKind.VALUES_CHANGED],
                    value_changes=[
                        ValueChange(index=i, phys_before=float(i), phys_after=float(i + 1))
                        for i in range(128)
                    ],
                )
            ]
        )
        target = tmp_path / "map.xlsx"
        ExcelDiffExporter().export(cs, target, {}, CTX)
        ws = load_workbook(target)["改动明细"]
        text = ws.cell(row=2, column=8).value
        assert "共 128 项" in text
        assert text.count("\n") == 20  # 20 行元素 + 1 行截断说明 - 1

    def test_wrong_payload_type_rejected(self, tmp_path):
        with pytest.raises(TypeError, match="ChangeSet"):
            ExcelDiffExporter().export(["不是 ChangeSet"], tmp_path / "x.xlsx", {}, CTX)

    def test_extension_is_enforced(self, tmp_path):
        target = tmp_path / "report.txt"
        result = ExcelDiffExporter().export(_changeset(), target, {}, CTX)
        assert result.suffix == ".xlsx"

    def test_plugin_registered_as_exporter(self):
        reg = PluginRegistry()
        reg.load_all([PLUGINS_DIR])
        ids = {p.manifest.id for p in reg}
        assert "export-excel-diff" in ids
        from automotive_dev_support.core.skeleton.capability import PluginCapability

        assert "export-excel-diff" in reg.summary()["capabilities"]["exporter"]
