"""标定比对报告的 Excel 导出器。

## 报告为什么是三张表

一张表打天下是标定报告最常见的事故。三类信息的使用场景
完全不同，混在一起谁都用不好：

    汇总      给不看细节的人（领导、归档）—— 几个数说清「改了多少」
    改动明细  给干活的人 —— 逐条可追溯，可直接贴进变更说明
    地址漂移  给怀疑人生的人 —— App 重编后地址全变了，
              单独一张表，证明「变的是编译产物，不是标定」

## 几个导出层面的决定

1. **数值变化按物理值展示。** 原始值是 ECU 里的整数，
   工程师认物理值。两者都在时以物理值为准。
2. **元素变化有上限。** 一个 16x8 的 MAP 全量逐元素列出来
   没人看，默认展示前 20 项并注明总数。
3. **地址一律 16 进制。** 标定工程里没人用十进制地址对话。
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from automotive_dev_support.core.models.calibration import (
    ChangeKind,
    ChangeSet,
    CharacteristicChange,
    ValueChange,
)
from automotive_dev_support.core.ports.plugin import Exporter
from automotive_dev_support.core.skeleton import JobContext

__all__ = ["ExcelDiffExporter"]

#: 单个量的元素变化最多展示多少项
_MAX_ELEMENTS = 20

_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_KIND_LABELS = {
    ChangeKind.ADDED: "新增",
    ChangeKind.REMOVED: "删除",
    ChangeKind.VALUES_CHANGED: "数值变化",
    ChangeKind.ATTRS_CHANGED: "属性变化",
    ChangeKind.ADDRESS_MOVED: "地址漂移",
}


class ExcelDiffExporter:
    """`core.ports.Exporter` 的实现。"""

    media_type = _MEDIA_TYPE
    extension = "xlsx"

    def export(
        self,
        payload: Any,
        target: Path,
        params: dict[str, Any],
        ctx: JobContext,
    ) -> Path:
        if not isinstance(payload, ChangeSet):
            raise TypeError(
                f"export-excel-diff 只接受 ChangeSet，收到 {type(payload).__name__}。"
                "请先运行比对（diff-calibration）再导出。"
            )

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover - 取决于安装环境
            raise RuntimeError(
                "导出 Excel 需要 openpyxl（pip install openpyxl）。"
                "它是可选依赖 —— 不装不影响其它功能，只是不能导出报告。"
            ) from exc

        max_elements = int(params.get("max_elements", _MAX_ELEMENTS))
        wb = Workbook()

        self._write_summary(wb, payload, Font, Alignment, get_column_letter)
        self._write_details(wb, payload, Font, Alignment, get_column_letter, max_elements)
        self._write_address_moves(wb, payload, Font, Alignment, get_column_letter)

        target = Path(target)
        if target.suffix != f".{self.extension}":
            target = target.with_suffix(f".{self.extension}")
        wb.save(target)
        return target

    # -- 表一：汇总 ----------------------------------------------------------

    def _write_summary(self, wb, cs: ChangeSet, Font, Alignment, get_column_letter) -> None:
        ws = wb.active
        ws.title = "汇总"
        s = cs.summary()

        rows: list[tuple[str, Any]] = [
            ("基准版本", cs.baseline_label or "（未命名）"),
            ("对比版本", cs.candidate_label or "（未命名）"),
            ("生成时间", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
            ("", ""),
            ("新增标定量", s["added"]),
            ("删除标定量", s["removed"]),
            ("改动标定量", s["changed"]),
            ("其中：纯地址漂移", s["address_moved"]),
            ("需要关注的改动", s["meaningful"]),
            ("数值变化条数", s["value_changes"]),
        ]
        for label, value in rows:
            ws.append([label, value])

        ws["A1"].font = Font(bold=True, size=12)
        for row in ws.iter_rows(min_row=1, max_row=len(rows), max_col=1):
            row[0].font = Font(bold=True)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 48

        note_row = len(rows) + 2
        ws.cell(row=note_row, column=1, value="说明").font = Font(bold=True)
        ws.cell(
            row=note_row + 1,
            column=1,
            value=(
                "「纯地址漂移」是 App 重编导致的地址变化，语义未变，"
                "默认不计入「需要关注的改动」。"
            ),
        )
        ws.cell(row=note_row + 1, column=1).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=2)

        if cs.warnings:
            warn_row = note_row + 3
            ws.cell(row=warn_row, column=1, value="警告").font = Font(bold=True)
            for i, warning in enumerate(cs.warnings, start=1):
                ws.cell(row=warn_row + i, column=1, value=f"{i}. {warning}")
                ws.merge_cells(
                    start_row=warn_row + i, start_column=1, end_row=warn_row + i, end_column=2
                )

    # -- 表二：改动明细 ------------------------------------------------------

    def _write_details(
        self, wb, cs: ChangeSet, Font, Alignment, get_column_letter, max_elements: int
    ) -> None:
        ws = wb.create_sheet("改动明细")
        headers = [
            "符号路径", "名称", "类别", "变化类型", "分组",
            "地址（前 → 后）", "属性变化", "数值变化",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        entries: list[tuple[str, CharacteristicChange]] = []
        entries.extend(("新增", c) for c in cs.added)
        entries.extend(("删除", c) for c in cs.removed)
        entries.extend(("改动", c) for c in cs.changed)
        # 纯地址漂移的放最后 —— 它们通常是噪音
        entries.sort(key=lambda item: 0 if item[1].is_meaningful else 1)

        for category, c in entries:
            ws.append([
                c.path,
                c.name,
                category,
                " / ".join(_KIND_LABELS.get(k, str(k)) for k in c.kinds),
                c.group,
                _address_text(c),
                _attrs_text(c),
                _values_text(c, max_elements),
            ])

        widths = [36, 20, 8, 22, 14, 30, 44, 60]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # -- 表三：地址漂移 ------------------------------------------------------

    def _write_address_moves(self, wb, cs: ChangeSet, Font, Alignment, get_column_letter) -> None:
        ws = wb.create_sheet("地址漂移")
        ws.append(["符号路径", "原地址", "新地址"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        moves = cs.address_moves()
        for c in sorted(moves, key=lambda c: c.path):
            ws.append([c.path, _hex(c.address_before), _hex(c.address_after)])

        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 16
        ws.freeze_panes = "A2"
        if not moves:
            ws.append(["（两个版本的地址完全一致）"])


def _address_text(c: CharacteristicChange) -> str:
    if c.address_before is None and c.address_after is None:
        return ""
    if c.address_before == c.address_after:
        return _hex(c.address_after)
    return f"{_hex(c.address_before)} → {_hex(c.address_after)}"


def _attrs_text(c: CharacteristicChange) -> str:
    if not c.attribute_changes:
        return ""
    lines = [fc.display for fc in c.attribute_changes]
    return "\n".join(lines)


def _values_text(c: CharacteristicChange, max_elements: int) -> str:
    if not c.value_changes:
        return ""
    shown = c.value_changes[:max_elements]
    lines = [_value_line(v) for v in shown]
    if len(c.value_changes) > len(shown):
        lines.append(f"……共 {len(c.value_changes)} 项，仅展示前 {len(shown)} 项")
    return "\n".join(lines)


def _value_line(v: ValueChange) -> str:
    where = f"[{v.index}]"
    if v.row is not None and v.col is not None:
        where = f"[{v.row},{v.col}]"
    before = "—" if v.phys_before is None else _num(v.phys_before)
    after = "—" if v.phys_after is None else _num(v.phys_after)
    return f"{where} {before} → {after}"


def _num(value: float) -> str:
    if value == int(value):
        return str(int(value))
    return f"{value:g}"


def _hex(address: int | None) -> str:
    return "—" if address is None else f"0x{address:X}"
